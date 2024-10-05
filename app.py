import os
import re
import docx
from shutil import move
import time
import docx2txt
from PyPDF2 import PdfReader
# from config import SECRET_KEY
from openpyxl import Workbook
from flask import Flask, render_template, request, flash, send_file
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
# app.secret_key = SECRET_KEY

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in {'docx', 'doc', 'pdf'}

def extract_contact_info(text):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    phone_pattern = r'\b(?:\+?(\d{1,3}))?[-. (]*(\d{3})[-. )]*(\d{3})[-. ]*(\d{4})\b'

    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)

    return emails, [f"{phone[0]}-{phone[1]}-{phone[2]}-{phone[3]}" for phone in phones]

def extract_text_from_cv(file_path):
    if file_path.endswith('.docx'):
        return docx2txt.process(file_path)
    elif file_path.endswith('.doc'):
        doc = docx.Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])
    elif file_path.endswith('.pdf'):
        text = ""
        with open(file_path, 'rb') as f:
            reader = PdfReader(f)
            for page in reader.pages:
                text += page.extract_text()
        return text
    else:
        raise ValueError("Unsupported file format. Only .doc, .docx, and .pdf files are supported.")

def write_to_excel(data, file_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["File Name", "Email", "Phone Number"])

    for entry in data:
        ws.append([entry["file_name"], entry["email"], entry["phone_number"]])

    wb.save(file_path)


@app.route('/', methods=['GET', 'POST'])
def upload_file():
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('No file part')
            return render_template('index.html')
        files = request.files.getlist('file')
        if len(files) == 0:
            flash('No selected file')
            return render_template('index.html')
        
        timestamp = int(time.time())
        folder_name = f"temp{timestamp}"
        folder_path = os.path.join(app.config['UPLOAD_FOLDER'], folder_name)
        os.makedirs(folder_path, exist_ok=True)
        
        for file in files:
            if file and allowed_file(file.filename):
                file_path = os.path.join(folder_path, secure_filename(file.filename))
                file.save(file_path)

        cv_data = []

        for filename in os.listdir(folder_path):
            file_path = os.path.join(folder_path, filename)
            text = extract_text_from_cv(file_path)
            emails, phones = extract_contact_info(text)
            cv_data.append({
                "file_name": filename,
                "email": ', '.join(emails),
                "phone_number": ', '.join(phones)
            })

        output_excel = os.path.join(os.path.expanduser('~'), 'Downloads', 'output.xlsx')
        write_to_excel(cv_data, output_excel)
        move(output_excel, os.path.join(os.path.expanduser('~'), 'Downloads', 'output.xlsx'))
        return send_file(os.path.join(os.path.expanduser('~'), 'Downloads', 'output.xlsx'), as_attachment=True)

    return render_template('index.html')


if __name__ == '__main__':
    app.run(debug=True)
